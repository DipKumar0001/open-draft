import sys
import json
import os

def out(content="", metadata={}, error=None):
    print(json.dumps({"content": content, "metadata": metadata, "error": error}))
    sys.exit(0)

def main():
    if len(sys.argv) < 2:
        out(error="No file path provided")
    
    filepath = sys.argv[1]
    if not os.path.exists(filepath):
        out(error=f"File not found: {filepath}")

    ext = os.path.splitext(filepath)[1].lower()
    content = ""
    metadata = {"extension": ext, "method": "python"}

    try:
        if ext == '.pdf':
            try:
                import fitz
                doc = fitz.open(filepath)
                content = chr(10).join([page.get_text() for page in doc])
            except ImportError:
                out(error="PyMuPDF (fitz) not installed")
                
        elif ext in ['.docx', '.doc']:
            try:
                import docx
                doc = docx.Document(filepath)
                content = chr(10).join([p.text for p in doc.paragraphs])
            except ImportError:
                out(error="python-docx not installed")
                
        elif ext in ['.xlsx', '.xls']:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(filepath, data_only=True)
                for sheet in wb.sheetnames:
                    ws = wb[sheet]
                    content += f"--- {sheet} ---\
"
                    for row in ws.iter_rows(values_only=True):
                        content += ",".join([str(cell) if cell is not None else "" for cell in row]) + "\
"
            except ImportError:
                out(error="openpyxl not installed")
                
        elif ext in ['.png', '.jpg', '.jpeg', '.tiff', '.bmp', '.gif', '.webp']:
            try:
                import pytesseract
                from PIL import Image
                content = pytesseract.image_to_string(Image.open(filepath))
            except ImportError:
                out(error="pytesseract or Pillow not installed")
                
        else:
            with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()

        out(content=content, metadata=metadata)
        
    except Exception as e:
        out(error=str(e))

if __name__ == "__main__":
    main()
